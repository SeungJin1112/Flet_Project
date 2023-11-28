import sqlite3 as sql3
import os
import openpyxl
import json

from context import *

g_sql3_instance = None;
##################################################
g_DEFINE_CREATE_TABLE_APPROXIMATELY = """
    CREATE TABLE IF NOT EXISTS T_APPROXIMATELY
    (
        seq INTEGER PRIMARY KEY AUTOINCREMENT,
        item_standard_code TEXT,
        product_name TEXT,
        product_name_eng TEXT,
        company_name TEXT,
        company_name_eng TEXT,
        main_ingredient TEXT,
        main_ingredient_eng TEXT,
        additives TEXT
    )
""";
#-------------------------------------------------
g_DEFINE_CREATE_TABLE_REPORT = """
""";
##################################################
class DbSqlite():
    _con = None;
    _c = None;

    def __init__(self):
        global g_sql3_instance;

        if g_sql3_instance == None:
            g_sql3_instance = self; 
    
        self.fn_init_create_tables();
        self.fn_init_insert_tables();
    
    def fn_start(self): pass;
    def fn_end(self): 
        if self._con != None: 
            self._con.close();

    def fn_enable(self): pass;
    def fn_disable(self): pass;
    def fn_get_instance(self): 
        if g_sql3_instance != None:
            return g_sql3_instance;
##################################################
    def fn_init_create_tables(self):
        self._con = sql3.connect('auto_medic.db');
        self._c = self._con.cursor();

        self._c.execute(g_DEFINE_CREATE_TABLE_APPROXIMATELY);
        self._c.execute(g_DEFINE_CREATE_TABLE_REPORT);

        self._con.commit();
        self._con.close();
##################################################
    def fn_init_insert_tables(self):
        insert_data_info = [
            {"file_name": "의약품등제품정보목록.xlsx", "table_name": "T_APPROXIMATELY", "placeholders": "(NULL,?,?,?,?,?,?,?,?)"}
        ]
        
        for data_info in insert_data_info:
            file_name = data_info["file_name"]
            table_name = data_info["table_name"]
            placeholders = data_info["placeholders"]

            if not os.path.exists(file_name):
                continue

            if self.fn_is_data_exist(table_name):
                continue

            try:
                workbook = openpyxl.load_workbook(file_name)
                sheet = workbook.active

                with sql3.connect('auto_medic.db') as con:
                    cursor = con.cursor()

                    for row in sheet.iter_rows(min_row=2):
                        data = (
                            row[0].value,  # 품목기준코드
                            row[1].value,  # 제품명
                            row[2].value,  # 제품영문명
                            row[3].value,  # 업체명
                            row[4].value,  # 업체영문명
                            row[10].value,  # 주성분
                            row[29].value,  # 주성분영문
                            row[11].value,  # 첨가제
                        )
                        cursor.execute(f"INSERT INTO {table_name} VALUES {placeholders}", data)
                    con.commit()

            except sql3.DatabaseError as e:
                print(e)
            except Exception as e:
                print(e)

    def fn_is_data_exist(self, table_name):
        self._con = sql3.connect('auto_medic.db')
        self._c = self._con.cursor()

        self._c.execute(f"SELECT COUNT(*) FROM {table_name}")
        count = self._c.fetchone()[0]
        self._con.close()

        return count > 0
##################################################
    def fn_json_to_sql_data(self, str_json): 
        # {"예상 병명": ["편두통", "감기"], "성분": ["아세트아미노펜", "진통제"]}
        data = json.loads(str_json);
        return data;

    def fn_get_main_ingredient(self, str_json):

        if str_json == None or str_json == "":
            return None;

        self._con = sql3.connect('auto_medic.db');
        self._c = self._con.cursor();

        query = [];
        main_ingredient = json.loads(str_json)["성분"];

        for i in main_ingredient:
            for row in self._c.execute(f"SELECT product_name FROM T_APPROXIMATELY WHERE main_ingredient LIKE '%{i}%'"):
                query.append(row[0]);
        
        self._con.close();

        return query;


